"""Data record storage service for Base Camp OS."""

from __future__ import annotations

import csv
import io
import json
import uuid

from sqlalchemy import and_, func, select

from app.basecamp.enums import FileFormat
from app.basecamp.schemas import (
    DataExportRequest,
    DataQueryRequest,
    DataQueryResponse,
    DataRecordCreate,
    DataRecordRead,
    DataRecordUpdate,
)
from app.db.models import DataRecord
from app.exceptions import TracecatNotFoundError
from app.service import BaseWorkspaceService


class DataRecordService(BaseWorkspaceService):
    """Service for managing data records."""

    service_name = "basecamp_storage"

    async def list_records(
        self,
        *,
        schema_id: uuid.UUID | None = None,
        job_id: uuid.UUID | None = None,
        limit: int = 100,
        offset: int = 0,
    ) -> list[DataRecordRead]:
        """List data records.

        Args:
            schema_id: Optional filter by schema.
            job_id: Optional filter by ingestion job.
            limit: Maximum number of results.
            offset: Number of results to skip.

        Returns:
            List of data records.
        """
        conditions = [DataRecord.workspace_id == self.workspace_id]

        if schema_id:
            conditions.append(DataRecord.schema_id == schema_id)
        if job_id:
            conditions.append(DataRecord.job_id == job_id)

        stmt = (
            select(DataRecord)
            .where(and_(*conditions))
            .order_by(DataRecord.created_at.desc())
            .limit(limit)
            .offset(offset)
        )

        result = await self.session.execute(stmt)
        records = result.scalars().all()

        return [self._record_to_read(r) for r in records]

    async def get_record(self, record_id: uuid.UUID) -> DataRecordRead:
        """Get a data record by ID.

        Args:
            record_id: The record ID.

        Returns:
            The data record.

        Raises:
            TracecatNotFoundError: If record not found.
        """
        record = await self._get_record_model(record_id)
        return self._record_to_read(record)

    async def create_record(self, params: DataRecordCreate) -> DataRecordRead:
        """Create a new data record.

        Args:
            params: The record creation parameters.

        Returns:
            The created record.
        """
        record = DataRecord(
            workspace_id=self.workspace_id,
            schema_id=params.schema_id,
            job_id=None,
            data=params.data,
        )

        self.session.add(record)
        await self.session.flush()
        await self.session.refresh(record)

        self.logger.info("Created data record", record_id=str(record.id))

        return self._record_to_read(record)

    async def update_record(
        self,
        record_id: uuid.UUID,
        params: DataRecordUpdate,
    ) -> DataRecordRead:
        """Update a data record.

        Args:
            record_id: The record ID.
            params: The update parameters.

        Returns:
            The updated record.

        Raises:
            TracecatNotFoundError: If record not found.
        """
        record = await self._get_record_model(record_id)
        record.data = params.data

        await self.session.flush()
        await self.session.refresh(record)

        self.logger.info("Updated data record", record_id=str(record_id))

        return self._record_to_read(record)

    async def delete_record(self, record_id: uuid.UUID) -> None:
        """Delete a data record.

        Args:
            record_id: The record ID.

        Raises:
            TracecatNotFoundError: If record not found.
        """
        record = await self._get_record_model(record_id)

        await self.session.delete(record)
        await self.session.flush()

        self.logger.info("Deleted data record", record_id=str(record_id))

    async def query(self, request: DataQueryRequest) -> DataQueryResponse:
        """Query data records with filtering.

        Args:
            request: The query request.

        Returns:
            The query response with matching records.
        """
        conditions = [
            DataRecord.workspace_id == self.workspace_id,
            DataRecord.schema_id == request.schema_id,
        ]

        # Apply JSONB filters
        for field, value in request.filters.items():
            if isinstance(value, list):
                # IN clause for lists
                conditions.append(
                    DataRecord.data[field].astext.in_([str(v) for v in value])
                )
            elif isinstance(value, dict):
                # Handle operators like {"$gt": 10}
                for op, op_value in value.items():
                    if op == "$eq":
                        conditions.append(
                            DataRecord.data[field].astext == str(op_value)
                        )
                    elif op == "$ne":
                        conditions.append(
                            DataRecord.data[field].astext != str(op_value)
                        )
                    elif op == "$gt":
                        conditions.append(
                            DataRecord.data[field].astext.cast(type_=None)
                            > str(op_value)
                        )
                    elif op == "$gte":
                        conditions.append(
                            DataRecord.data[field].astext.cast(type_=None)
                            >= str(op_value)
                        )
                    elif op == "$lt":
                        conditions.append(
                            DataRecord.data[field].astext.cast(type_=None)
                            < str(op_value)
                        )
                    elif op == "$lte":
                        conditions.append(
                            DataRecord.data[field].astext.cast(type_=None)
                            <= str(op_value)
                        )
                    elif op == "$like":
                        conditions.append(
                            DataRecord.data[field].astext.ilike(f"%{op_value}%")
                        )
                    elif op == "$contains":
                        conditions.append(DataRecord.data[field].contains(op_value))
            else:
                # Exact match
                conditions.append(DataRecord.data[field].astext == str(value))

        # Count total
        count_stmt = select(func.count()).select_from(
            select(DataRecord.id).where(and_(*conditions)).subquery()
        )
        count_result = await self.session.execute(count_stmt)
        total_count = count_result.scalar() or 0

        # Build query
        stmt = select(DataRecord).where(and_(*conditions))

        # Apply ordering
        if request.order_by:
            order_col = DataRecord.data[request.order_by].astext
            if request.order_direction == "asc":
                stmt = stmt.order_by(order_col.asc())
            else:
                stmt = stmt.order_by(order_col.desc())
        else:
            stmt = stmt.order_by(DataRecord.created_at.desc())

        # Apply pagination
        stmt = stmt.limit(request.limit).offset(request.offset)

        result = await self.session.execute(stmt)
        records = result.scalars().all()

        return DataQueryResponse(
            records=[self._record_to_read(r) for r in records],
            total_count=total_count,
            limit=request.limit,
            offset=request.offset,
        )

    async def export(self, request: DataExportRequest) -> bytes:
        """Export data records to a file format.

        Args:
            request: The export request.

        Returns:
            The exported data as bytes.
        """
        # Get all matching records
        query_request = DataQueryRequest(
            schema_id=request.schema_id,
            filters=request.filters,
            limit=10000,  # Reasonable export limit
            offset=0,
        )
        response = await self.query(query_request)

        if request.format == FileFormat.JSON:
            return self._export_json(response.records)
        elif request.format == FileFormat.CSV:
            return self._export_csv(response.records)
        elif request.format == FileFormat.EXCEL:
            return self._export_excel(response.records)
        else:
            return self._export_json(response.records)

    def _export_json(self, records: list[DataRecordRead]) -> bytes:
        """Export records as JSON.

        Args:
            records: The records to export.

        Returns:
            JSON bytes.
        """
        data = [r.data for r in records]
        return json.dumps(data, indent=2, default=str).encode("utf-8")

    def _export_csv(self, records: list[DataRecordRead]) -> bytes:
        """Export records as CSV.

        Args:
            records: The records to export.

        Returns:
            CSV bytes.
        """
        if not records:
            return b""

        # Collect all field names
        all_fields: set[str] = set()
        for record in records:
            all_fields.update(record.data.keys())

        fields = sorted(all_fields)

        output = io.StringIO()
        writer = csv.DictWriter(output, fieldnames=fields)
        writer.writeheader()

        for record in records:
            row = {field: record.data.get(field, "") for field in fields}
            writer.writerow(row)

        return output.getvalue().encode("utf-8")

    def _export_excel(self, records: list[DataRecordRead]) -> bytes:
        """Export records as Excel.

        Args:
            records: The records to export.

        Returns:
            Excel bytes.
        """
        try:
            import openpyxl
        except ImportError:
            # Fall back to CSV if openpyxl not available
            return self._export_csv(records)

        if not records:
            wb = openpyxl.Workbook()
            output = io.BytesIO()
            wb.save(output)
            return output.getvalue()

        # Collect all field names
        all_fields: set[str] = set()
        for record in records:
            all_fields.update(record.data.keys())

        fields = sorted(all_fields)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Data"

        # Write header
        for col, field in enumerate(fields, start=1):
            ws.cell(row=1, column=col, value=field)

        # Write data
        for row_idx, record in enumerate(records, start=2):
            for col_idx, field in enumerate(fields, start=1):
                value = record.data.get(field)
                if value is not None:
                    ws.cell(row=row_idx, column=col_idx, value=str(value))

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    async def _get_record_model(self, record_id: uuid.UUID) -> DataRecord:
        """Get the raw record model.

        Args:
            record_id: The record ID.

        Returns:
            The record model.

        Raises:
            TracecatNotFoundError: If not found.
        """
        stmt = select(DataRecord).where(
            and_(
                DataRecord.workspace_id == self.workspace_id,
                DataRecord.id == record_id,
            )
        )

        result = await self.session.execute(stmt)
        record = result.scalar_one_or_none()

        if not record:
            raise TracecatNotFoundError(f"Data record {record_id} not found")

        return record

    def _record_to_read(self, record: DataRecord) -> DataRecordRead:
        """Convert record model to read schema.

        Args:
            record: The record model.

        Returns:
            The read schema.
        """
        return DataRecordRead(
            id=record.id,
            schema_id=record.schema_id,
            job_id=record.job_id,
            data=record.data,
            created_at=record.created_at,
            updated_at=record.updated_at,
        )
